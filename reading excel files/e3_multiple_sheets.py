import openpyxl

# Load an Excel workbook
wb = openpyxl.load_workbook("excels/pcmall_sales.xlsx")

# list all sheets in the workbook
print("All sheets are: ", wb.sheetnames)
print()

# Get the first sheet
sheet = wb["2024-02"]
print("\n Sheet title: ", sheet.title)

# Read data from the first few rows
for row in sheet.iter_rows(min_row=1, max_row=5, values_only=True):
    print(row)  # Prints each row as a tuple

# # Get the second sheet
# sheet = wb["Sales Data B"]
# print("\n Sheet title: ", sheet.title)

# # Read data from the first few rows
# for row in sheet.iter_rows(min_row=1, max_row=5, values_only=True):
#     print(row)  # Prints each row as a tuple