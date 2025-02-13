import openpyxl

# Load an Excel workbook
wb = openpyxl.load_workbook("excels/pcmall_sales.xlsx")
sheet_list = wb.sheetnames

monthly_sales_list = []

# Monthly sales summary
for sheet_name in sheet_list:
    print("\n Sales month: ", sheet_name)
    sheet = wb[sheet_name]
    monthly_sales = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        # TO DO: Get the sales amount
        sales = row[1]
        # TO DO: Add sales to monthly
        monthly_sales += sales

    print(f"Monthly sales for {sheet_name}: {monthly_sales}")

    # Add the monthly sales to the list
    monthly_sales_list.append([sheet_name, monthly_sales])

# Create a new sheet to store the monthly sales summary
summary_sheet = wb.create_sheet(title="Summary")
for row in monthly_sales_list:
    # TO DO: Add the row to the summary sheet
    summary_sheet.append(row)

wb.save("outputs/pcmall_sales.xlsx")
print("Excel file created successfully.")
