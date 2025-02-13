import openpyxl

# Load an Excel workbook
wb = openpyxl.load_workbook("excels/pcmall_sales.xlsx")

# list all sheets in the workbook
print("All sheets are: ", wb.sheetnames)
print()

sheet_list = wb.sheetnames

# Monthly sales summary
for sheet_name in sheet_list:
    print("\n Sales month: ", sheet_name)
    sheet = wb[sheet_name]
    monthly_sales = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        # TO DO: Get the sales amount
        sales = row[1]

        # TO DO: Add sales to monthly
        monthly_sales += sales
    
    print(f"Monthly sales for {sheet_name}: {monthly_sales}")
