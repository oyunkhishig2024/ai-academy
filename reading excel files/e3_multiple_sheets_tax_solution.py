import openpyxl

# Load an Excel workbook
wb = openpyxl.load_workbook("excels/pcmall_sales.xlsx")

# list all sheets in the workbook
print("All sheets are: ", wb.sheetnames)
print()

sheet_list = wb.sheetnames

# Tax 10% then calculate monthly sales without tax
for sheet_name in sheet_list:
    print("\n Sales month: ", sheet_name)
    sheet = wb[sheet_name]
    monthly_sales = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        # TO DO: Get the sales amount
        sales = row[1]

        # TO DO: Calculate the tax
        tax = sales * 0.1

        # TO DO: Calculate the monthly sales without tax
        sales_without_tax = sales - tax

        # TO DO: Add sales to monthly_sales
        monthly_sales += sales_without_tax

    print(f"Monthly sales for {sheet_name} without tax: {monthly_sales}")
